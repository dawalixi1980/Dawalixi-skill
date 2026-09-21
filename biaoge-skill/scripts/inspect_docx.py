"""诊断 Word 文档的表格与版式环境。

用法:
    python inspect_docx.py <file.docx>
"""
import io
import json
import os
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.oxml.ns import qn

import wtlib as L


def main():
    if len(sys.argv) < 2:
        print('用法: python inspect_docx.py <file.docx>')
        sys.exit(1)
    path = sys.argv[1]
    doc = Document(path)

    print('=' * 62)
    print('文档:', os.path.basename(path))
    print('=' * 62)

    # 分节 / 分栏
    secs = L.body_sections(doc)
    print('\n[分节与分栏]')
    for i, (s, e, sp) in enumerate(secs, 1):
        avail, num, content = L.section_avail(sp)
        pg = sp.find(qn('w:pgSz'))
        w = int(pg.get(qn('w:w')))
        h = int(pg.get(qn('w:h')))
        orient = pg.get(qn('w:orient')) or 'portrait'
        print('  第%d节: 页面 %.0f×%.0fmm  %s  版心 %.0fmm  栏数 %d  单栏宽 %.1fmm'
              % (i, w / 56.6929, h / 56.6929, orient, content / 56.6929, num, avail / 56.6929))

    # 页眉页脚残留
    print('\n[页眉/页脚检查]')
    import zipfile
    z = zipfile.ZipFile(path)
    flags = []
    for n in z.namelist():
        if 'header' in n and n.endswith('.xml'):
            x = z.read(n).decode('utf-8', 'ignore')
            txt = ''.join(t for t in __import__('re').findall(r'<w:t[^>]*>([^<]*)</w:t>', x) if t.strip())
            imgs = 'graphicData' in x or 'v:imagedata' in x
            if txt or imgs:
                flags.append((n, txt[:80], imgs))
        if 'footer' in n and n.endswith('.xml'):
            x = z.read(n).decode('utf-8', 'ignore')
            txt = ''.join(t for t in __import__('re').findall(r'<w:t[^>]*>([^<]*)</w:t>', x) if t.strip())
            imgs = 'graphicData' in x or 'v:imagedata' in x
            if txt or imgs:
                flags.append((n, txt[:80], imgs))
    if flags:
        for n, t, im in flags:
            print('  ⚠ %s  文本="%s"  图片=%s' % (n, t, im))
        print('  → 若含与本项目无关的名称/图签，请用 fit_tables.py --clean-header 清除')
    else:
        print('  干净（无文字、无图片）')

    # 表格
    tmap = L.table_section_map(doc)
    from collections import Counter
    cnt = Counter(len(t.columns) for t in doc.tables)
    print('\n[表格概况]  共 %d 张' % len(doc.tables))
    print('  列数分布:', dict(sorted(cnt.items())))

    print('\n[逐表列宽与内容余量]')
    print('  %-4s %-4s %-9s %-9s %s' % ('表', '列数', '当前总宽', '可用栏宽', '判定'))
    over = 0
    for i, t in enumerate(doc.tables):
        avail = tmap.get(i, (0, 1, 0))[0]
        cur = sum(int(c.width) for c in t.columns if c.width is not None) / 635.0
        verdict = '偏窄(可放宽)' if cur < avail * 0.9 else ('超宽!' if cur > avail + 20 else '合适')
        if cur > avail + 20:
            over += 1
        print('  %-4d %-4d %-9.1f %-9.1f %s'
              % (i + 1, len(t.columns), cur / 56.6929, avail / 56.6929, verdict))
    if over:
        print('  ⚠ %d 张表超出可用栏宽（多栏文档最常见问题）' % over)


if __name__ == '__main__':
    main()
