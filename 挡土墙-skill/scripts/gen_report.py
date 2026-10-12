# -*- coding: utf-8 -*-
"""
计算书 Word 生成器
==================
读取输入 JSON，调用 dantu_calc 计算，输出理正风格的 Word 计算书（.docx）。

用法：
    python gen_report.py input.json 输出.docx [--cover] [--unit 单位名] [--stage 施工图]

依赖：python-docx
"""
import sys
import os
import argparse

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import dantu_calc  # noqa: E402

from docx import Document  # noqa: E402
from docx.shared import Pt, Cm  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402


def _set_cn_font(run, name="宋体", size=None, bold=False):
    run.font.name = name
    run.font.bold = bold
    if size:
        run.font.size = Pt(size)
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        from docx.oxml import OxmlElement
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)


def add_cover(doc, project, unit, stage):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_cn_font(p.add_run(project), "黑体", 22, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_cn_font(p.add_run("挡土墙设计计算书"), "黑体", 22, True)
    for _ in range(4):
        doc.add_paragraph()
    for label in ["设计阶段", "计 算 人", "校 对 人", "审 核 人"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_cn_font(p.add_run("%s：" % label), "宋体", 14)
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_cn_font(p.add_run(unit), "宋体", 16, True)
    doc.add_page_break()


def build_docx(inp, out_path, cover=False, unit="", stage="施工图"):
    r = dantu_calc.calc_gravity(inp)
    text = dantu_calc.report_text(inp, r)
    doc = Document()
    # 页面
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(2.5)

    if cover:
        add_cover(doc, inp.get("project_name", "挡土墙工程"), unit, stage)

    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.5
        if i == 0:
            _set_cn_font(p.add_run(line), "黑体", 14, True)
        else:
            _set_cn_font(p.add_run(line), "宋体", 10.5)
    doc.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("output")
    ap.add_argument("--cover", action="store_true")
    ap.add_argument("--unit", default="")
    ap.add_argument("--stage", default="施工图")
    a = ap.parse_args()
    import json
    with open(a.input, "r", encoding="utf-8-sig") as fp:
        inp = json.load(fp)
    build_docx(inp, a.output, a.cover, a.unit, a.stage)
    print("已生成:", a.output)


if __name__ == "__main__":
    main()
