# -*- coding: utf-8 -*-
"""
导出 .xls / .xlsx / .pdf 表格内容为文本，便于提取预算数据。

用法:
    python dump_tables.py <文件> [输出.txt]

- .xls/.xlsx : 逐工作表输出 "R<行>: [列]值 | [列]值"
- .pdf       : 逐页输出抽取文本
依赖: xlrd(.xls) / openpyxl(.xlsx) / pypdf(.pdf)   —— 按需安装
"""
import sys
import os
import io


def dump_xls(path):
    import xlrd
    bk = xlrd.open_workbook(path)
    out = ["SHEETS: %s" % bk.sheet_names()]
    for sh in bk.sheets():
        out.append("\n===== SHEET: %s (rows=%d cols=%d) =====" % (sh.name, sh.nrows, sh.ncols))
        for r in range(sh.nrows):
            cells = []
            for c in range(sh.ncols):
                v = sh.cell_value(r, c)
                if v not in ("", None):
                    cells.append("[%d]%s" % (c, v))
            if cells:
                out.append("R%d: " % r + " | ".join(cells))
    return out


def dump_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    out = ["SHEETS: %s" % wb.sheetnames]
    for ws in wb.worksheets:
        out.append("\n===== SHEET: %s (%dx%d) =====" % (ws.title, ws.max_row, ws.max_column))
        for r in range(1, ws.max_row + 1):
            cells = []
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if v not in ("", None):
                    cells.append("[%d]%s" % (c, v))
            if cells:
                out.append("R%d: " % r + " | ".join(cells))
    return out


def dump_pdf(path):
    from pypdf import PdfReader
    r = PdfReader(path)
    out = ["PAGES: %d" % len(r.pages)]
    for i, p in enumerate(r.pages):
        out.append("\n===== PAGE %d =====" % (i + 1))
        out.append(p.extract_text() or "")
    return out


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    path = sys.argv[1]
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls":
        text = dump_xls(path)
    elif ext in (".xlsx", ".xlsm"):
        text = dump_xlsx(path)
    elif ext == ".pdf":
        text = dump_pdf(path)
    else:
        sys.exit("不支持的格式: %s" % ext)
    data = "\n".join(text)
    if len(sys.argv) >= 3:
        with io.open(sys.argv[2], "w", encoding="utf-8") as f:
            f.write(data)
        print("已写出:", sys.argv[2])
    else:
        sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
        print(data)


if __name__ == "__main__":
    main()
